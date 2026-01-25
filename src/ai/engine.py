"""
AI categorization and NSFW scoring pipeline.
"""

from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

from config import AppConfig
from database import DatabaseManager
from utils import ResourceMonitor

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

try:
    from transformers import CLIPModel, CLIPProcessor
except ImportError:  # pragma: no cover
    CLIPModel = None
    CLIPProcessor = None

try:
    from nudenet import NudeClassifier
except ImportError:  # pragma: no cover
    NudeClassifier = None

try:
    from sentence_transformers import SentenceTransformer, util as st_util
except ImportError:  # pragma: no cover
    SentenceTransformer = None
    st_util = None

try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover
    fitz = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


IMAGE_EXTS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".bmp",
    ".tif",
    ".tiff",
    ".webp",
    ".heic",
    ".heif",
}
VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".wmv", ".flv", ".webm"}
DOC_EXTS = {
    ".pdf",
    ".doc",
    ".docx",
    ".ppt",
    ".pptx",
    ".xls",
    ".xlsx",
    ".odt",
    ".ods",
    ".odp",
    ".rtf",
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".epub",
    ".mobi",
    ".azw",
    ".azw3",
}
AUDIO_EXTS = {".mp3", ".flac", ".wav", ".aac", ".m4a", ".ogg", ".wma", ".alac", ".aiff", ".ape"}


@dataclass
class AiCategorizationStats:
    """Summary stats for AI categorization."""

    processed: int
    skipped: int
    errors: int
    classified: int
    nsfw_flagged: int


class AiCategorizationEngine:
    """Categorize files and compute NSFW scores."""

    def __init__(
        self,
        config: AppConfig,
        db_manager: DatabaseManager,
        logger: Optional[logging.Logger] = None,
        monitor: Optional[ResourceMonitor] = None,
    ) -> None:
        self.config = config
        self.db_manager = db_manager
        self.logger = logger or logging.getLogger("file_manager")
        self.monitor = monitor
        self.enabled = bool(self.config.get("ai", "enabled", default=True))
        self.reprocess_existing = bool(self.config.get("ai", "reprocess_existing", default=False))
        self.max_files = int(self.config.get("ai", "max_files_per_run", default=0))
        self.text_max_chars = int(self.config.get("ai", "text_max_chars", default=4000))
        self.image_batch_size = max(
            1,
            int(
            self.config.get(
                "ai",
                "image_batch_size",
                default=self.config.get("ai_models", "batch_size", default=8),
            )
        ),
        )
        self.use_embeddings = bool(self.config.get("ai", "use_embeddings", default=False))
        self.embedding_model_name = str(
            self.config.get("ai", "embedding_model", default="all-MiniLM-L6-v2")
        )
        self.classify_audio = bool(self.config.get("ai", "classify_audio", default=True))
        self.classify_unknown = bool(self.config.get("ai", "classify_unknown", default=False))
        self.nsfw_threshold = float(
            self.config.get(
                "ai", "nsfw_threshold", default=self.config.get("ai_models", "nsfw_threshold", default=0.85)
            )
        )
        self.image_categories = self.config.get(
            "ai",
            "image_categories",
            default=[
                "People",
                "Landscapes",
                "Documents",
                "Screenshots",
                "Memes",
                "Pets",
                "Food",
                "Travel",
                "Artwork",
            ],
        )
        self.document_categories = self.config.get(
            "ai",
            "document_categories",
            default={
                "Financial": ["invoice", "receipt", "bank", "tax", "balance", "statement", "payment"],
                "Medical": ["medical", "insurance", "doctor", "prescription", "clinic", "patient"],
                "Legal": ["contract", "agreement", "nda", "license", "legal", "court"],
                "Work": ["project", "meeting", "proposal", "report", "budget", "plan"],
                "Personal": ["resume", "cv", "passport", "id", "certificate", "letter"],
                "Education": ["course", "lecture", "assignment", "homework", "syllabus"],
            },
        )
        self._clip_model = None
        self._clip_processor = None
        self._nude_classifier = NudeClassifier() if NudeClassifier is not None else None
        self._embedding_model = None
        self._category_embeddings = None

    def run(self) -> Optional[AiCategorizationStats]:
        if not self.enabled:
            self.logger.info("AI categorization disabled.")
            return None

        processed = skipped = errors = classified = nsfw_flagged = 0
        image_batch: list[tuple] = []

        for entry in self.db_manager.iter_inventory(accessible_only=True):
            if self.max_files and processed >= self.max_files:
                break
            path = Path(entry.file_path)
            if not path.exists():
                skipped += 1
                continue
            if not self.reprocess_existing and self.db_manager.classification_exists(entry.file_id):
                skipped += 1
                continue
            if self.monitor is not None:
                self.monitor.throttle()

            ext = path.suffix.lower()
            if ext in IMAGE_EXTS and self.image_batch_size > 1:
                image_batch.append((entry, path))
                if len(image_batch) >= self.image_batch_size:
                    batch_stats = self._process_image_batch(image_batch)
                    processed += batch_stats[0]
                    classified += batch_stats[1]
                    nsfw_flagged += batch_stats[2]
                    errors += batch_stats[3]
                    image_batch = []
                continue

            processed += 1
            try:
                if ext in IMAGE_EXTS:
                    category, tags, confidence = self._classify_image(path)
                    nsfw_score = self._score_nsfw(path)
                    if nsfw_score is not None and nsfw_score >= self.nsfw_threshold:
                        nsfw_flagged += 1
                    self._save_classification(entry.file_id, category, "image", tags, nsfw_score, confidence)
                    classified += 1
                elif ext in DOC_EXTS:
                    category, tags, confidence = self._classify_document(path)
                    self._save_classification(entry.file_id, category, "document", tags, None, confidence)
                    classified += 1
                elif ext in VIDEO_EXTS:
                    category, tags, confidence = self._classify_video(path)
                    nsfw_score = self._score_nsfw(path)
                    if nsfw_score is not None and nsfw_score >= self.nsfw_threshold:
                        nsfw_flagged += 1
                    self._save_classification(entry.file_id, category, "video", tags, nsfw_score, confidence)
                    classified += 1
                elif ext in AUDIO_EXTS and self.classify_audio:
                    category, tags, confidence = self._classify_audio(path)
                    self._save_classification(entry.file_id, category, "audio", tags, None, confidence)
                    classified += 1
                elif self.classify_unknown:
                    category, tags, confidence = self._classify_unknown(path)
                    self._save_classification(entry.file_id, category, "other", tags, None, confidence)
                    classified += 1
                else:
                    skipped += 1
            except Exception as exc:
                errors += 1
                self.logger.warning("AI classification failed for %s: %s", path, exc)

        if image_batch:
            batch_stats = self._process_image_batch(image_batch)
            processed += batch_stats[0]
            classified += batch_stats[1]
            nsfw_flagged += batch_stats[2]
            errors += batch_stats[3]

        return AiCategorizationStats(
            processed=processed,
            skipped=skipped,
            errors=errors,
            classified=classified,
            nsfw_flagged=nsfw_flagged,
        )

    def _save_classification(
        self,
        file_id: int,
        category: str,
        subcategory: str,
        tags: list[str],
        nsfw_score: Optional[float],
        confidence: Optional[float],
    ) -> None:
        model = "clip" if self._clip_model is not None else "keyword"
        details = json.dumps({"tags": tags})
        self.db_manager.upsert_classification(
            file_id=file_id,
            category=category,
            subcategory=subcategory,
            tags=", ".join(tags),
            nsfw_score=nsfw_score,
            model=model,
            confidence=confidence,
            details_json=details,
        )

    def _classify_image(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        if Image is None:
            return self._fallback_label(path)
        self._ensure_clip_loaded()
        if self._clip_model is None or self._clip_processor is None:
            return self._fallback_label(path)
        try:
            image = Image.open(path).convert("RGB")
            inputs = self._clip_processor(
                text=self.image_categories, images=image, return_tensors="pt", padding=True
            )
            outputs = self._clip_model(**inputs)
            probs = outputs.logits_per_image.softmax(dim=1)
            best_idx = int(probs.argmax())
            confidence = float(probs[0, best_idx])
            label = str(self.image_categories[best_idx])
            return label, [label], confidence
        except Exception:
            return self._fallback_label(path)

    def _process_image_batch(self, batch: list[tuple]) -> tuple[int, int, int, int]:
        processed = len(batch)
        classified = nsfw_flagged = errors = 0
        paths = [path for _, path in batch]
        results = self._classify_images_batch(paths)
        for (entry, path), result in zip(batch, results):
            try:
                category, tags, confidence = result
                nsfw_score = self._score_nsfw(path)
                if nsfw_score is not None and nsfw_score >= self.nsfw_threshold:
                    nsfw_flagged += 1
                self._save_classification(entry.file_id, category, "image", tags, nsfw_score, confidence)
                classified += 1
            except Exception as exc:
                errors += 1
                self.logger.warning("AI image batch failed for %s: %s", path, exc)
        return processed, classified, nsfw_flagged, errors

    def _classify_images_batch(self, paths: list[Path]) -> list[tuple[str, list[str], Optional[float]]]:
        results = [self._fallback_label(path) for path in paths]
        if Image is None:
            return results
        self._ensure_clip_loaded()
        if self._clip_model is None or self._clip_processor is None:
            return results
        images = []
        indices = []
        opened = []
        for index, path in enumerate(paths):
            try:
                image = Image.open(path)
                opened.append(image)
                images.append(image.convert("RGB"))
                indices.append(index)
            except Exception:
                continue
        if not images:
            return results
        try:
            inputs = self._clip_processor(
                text=self.image_categories, images=images, return_tensors="pt", padding=True
            )
            outputs = self._clip_model(**inputs)
            probs = outputs.logits_per_image.softmax(dim=1)
            for offset, index in enumerate(indices):
                best_idx = int(probs[offset].argmax())
                confidence = float(probs[offset, best_idx])
                label = str(self.image_categories[best_idx])
                results[index] = (label, [label], confidence)
        except Exception:
            return results
        finally:
            for image in opened:
                try:
                    image.close()
                except Exception:
                    pass
        return results

    def _classify_video(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        return "Video", ["video"], None

    def _classify_audio(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        return "Audio", ["audio"], None

    def _classify_unknown(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        return "Uncategorized", ["Uncategorized"], 0.0

    def _classify_document(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        text = self._extract_text(path)
        if not text:
            return self._fallback_label(path)
        if self.use_embeddings:
            embedded = self._classify_document_embedding(text)
            if embedded is not None:
                return embedded
        scores = {}
        lowered = text.lower()
        for category, keywords in self.document_categories.items():
            hits = sum(1 for word in keywords if word.lower() in lowered)
            scores[category] = hits
        best_category = max(scores.items(), key=lambda item: item[1])[0] if scores else "Uncategorized"
        hits = scores.get(best_category, 0)
        confidence = float(hits) / max(sum(scores.values()), 1)
        tags = [best_category] if hits > 0 else ["Uncategorized"]
        return best_category, tags, confidence

    def _classify_document_embedding(self, text: str) -> Optional[tuple[str, list[str], Optional[float]]]:
        if SentenceTransformer is None or st_util is None:
            return None
        model = self._ensure_embedding_model()
        if model is None:
            return None
        categories = list(self.document_categories.keys())
        if not categories:
            return None
        if self._category_embeddings is None:
            self._category_embeddings = model.encode(categories, normalize_embeddings=True)
        text_embedding = model.encode([text], normalize_embeddings=True)
        scores = st_util.cos_sim(text_embedding, self._category_embeddings)[0]
        best_idx = int(scores.argmax())
        confidence = float(scores[best_idx])
        category = str(categories[best_idx])
        return category, [category], confidence

    def _extract_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf" and fitz is not None:
            try:
                doc = fitz.open(path)
                text = "".join(page.get_text() for page in doc[:2])
                doc.close()
                return text[: self.text_max_chars]
            except Exception:
                return ""
        if suffix == ".docx" and Document is not None:
            try:
                doc = Document(path)
                text = "\n".join(p.text for p in doc.paragraphs)
                return text[: self.text_max_chars]
            except Exception:
                return ""
        if suffix in {".xlsx", ".xls"} and openpyxl is not None:
            try:
                workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
                sheet = workbook.active
                values = []
                for row in sheet.iter_rows(min_row=1, max_row=20, max_col=10, values_only=True):
                    for value in row:
                        if value:
                            values.append(str(value))
                workbook.close()
                return " ".join(values)[: self.text_max_chars]
            except Exception:
                return ""
        if suffix in {".txt", ".md", ".csv", ".json", ".xml", ".rtf"}:
            try:
                return path.read_text(encoding="utf-8", errors="ignore")[: self.text_max_chars]
            except Exception:
                return ""
        return ""

    def _score_nsfw(self, path: Path) -> Optional[float]:
        if path.suffix.lower() not in IMAGE_EXTS:
            return None
        if self._nude_classifier is not None:
            try:
                scores = self._nude_classifier.classify(str(path))
                result = scores.get(str(path), {})
                return float(result.get("unsafe", result.get("nsfw", 0.0)))
            except Exception:
                return None
        return None

    def _fallback_label(self, path: Path) -> tuple[str, list[str], Optional[float]]:
        tokens = self._tokenize(path.stem)
        for label in self.image_categories:
            label_tokens = self._tokenize(label)
            if label_tokens.intersection(tokens):
                return label, [label], 0.1
        return "Uncategorized", ["Uncategorized"], 0.0

    def _tokenize(self, value: str) -> set[str]:
        return set(token for token in re.split(r"[^a-z0-9]+", value.lower()) if token)

    def _ensure_clip_loaded(self) -> None:
        if self._clip_model is not None or CLIPModel is None or CLIPProcessor is None:
            return
        model_name = self.config.get("ai", "clip_model", default=None)
        if not model_name:
            model_name = "openai/clip-vit-base-patch32"
        try:
            self._clip_model = CLIPModel.from_pretrained(model_name)
            self._clip_processor = CLIPProcessor.from_pretrained(model_name)
        except Exception:
            self._clip_model = None
            self._clip_processor = None

    def _ensure_embedding_model(self):
        if self._embedding_model is not None or SentenceTransformer is None:
            return self._embedding_model
        try:
            self._embedding_model = SentenceTransformer(self.embedding_model_name)
        except Exception:
            self._embedding_model = None
        return self._embedding_model
